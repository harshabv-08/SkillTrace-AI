import os
import re

from pypdf import PdfReader
from docx import Document


def clean_text(text):
    """Clean and normalize extracted resume text."""

    if not text:
        return ""

    text = text.replace("\x00", " ")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_pdf_text(file_path):
    """Extract text from a PDF file."""

    reader = PdfReader(file_path)
    pages = []

    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
            pages.append(page_text)
        except Exception:
            continue

    return clean_text("\n".join(pages))


def extract_docx_text(file_path):
    """Extract text from a DOCX file."""

    document = Document(file_path)
    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    # Also read simple table content.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))

    return clean_text("\n".join(paragraphs))


def extract_txt_text(file_path):
    """Extract text from a TXT file."""

    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        return clean_text(file.read())


def extract_resume_text(file_path):
    """Detect file type and extract resume text."""

    if not os.path.exists(file_path):
        raise FileNotFoundError("Resume file was not found.")

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    if extension == ".docx":
        return extract_docx_text(file_path)

    if extension == ".txt":
        return extract_txt_text(file_path)

    raise ValueError("Unsupported resume format. Use PDF, DOCX, or TXT.")